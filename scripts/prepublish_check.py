"""Reject private files before a public repository release without printing contents."""

from __future__ import annotations

from pathlib import Path
from datetime import date, datetime
import hashlib
import json
import sys
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


ALLOWED_WORKBOOKS = {
    Path("templates") / "秋招个人资料库模板.xlsx": "资料类型：空白模板",
    Path("examples") / "示例资料库.xlsx": "资料类型：虚构示例",
    Path("skill/recruitment-form-autofill/assets/秋招个人资料库模板.xlsx"): "资料类型：空白模板",
}
# Approved public cell contents, independent of the files being inspected.
# Formatting-only changes do not affect these fingerprints. Deliberate public
# fixture changes require reviewing the values before updating a fingerprint.
PUBLIC_CONTENT_HASHES = {
    "资料类型：空白模板": "9864463f08a1a9a99cc0b8251c6eb89427e521d584786515b93c6ee976812bb4",
    "资料类型：虚构示例": "af1efa8f2e2979c35f8ceaa2a514c452a289bb4ab67a143c35415beca958ff4f",
}
PROHIBITED_EXTENSIONS = {".docx", ".pdf", ".jpg", ".jpeg", ".png"}


def _iter_files(root: Path):
    for path in root.rglob("*"):
        if path.is_file() and not {".git", "node_modules", "__pycache__"}.intersection(path.parts):
            yield path


def _workbook_label_is_expected(path: Path, expected_label: str) -> bool:
    try:
        workbook = load_workbook(path, read_only=False, data_only=False)
        try:
            if "使用说明" not in workbook.sheetnames or workbook["使用说明"]["A2"].value != expected_label:
                return False
            contents = []
            for sheet in workbook:
                cells = []
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.comment is not None or cell.hyperlink is not None:
                            return False
                        if cell.value is not None:
                            value = cell.value.isoformat() if isinstance(cell.value, (date, datetime)) else cell.value
                            cells.append((cell.coordinate, cell.data_type, value))
                contents.append((sheet.title, cells))
            encoded = json.dumps(contents, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
            return hashlib.sha256(encoded).hexdigest() == PUBLIC_CONTENT_HASHES[expected_label]
        finally:
            workbook.close()
    except (OSError, ValueError, KeyError, BadZipFile, InvalidFileException, ParseError):
        return False


def check(root: Path) -> list[str]:
    root = Path(root).resolve()
    issues = []
    for path in _iter_files(root):
        relative = path.relative_to(root)
        if path.name == "profile-config.json":
            issues.append("local profile configuration")
        elif path.name.endswith(".personal.xlsx"):
            issues.append("personal workbook")
        elif path.suffix.lower() in PROHIBITED_EXTENSIONS:
            issues.append("personal attachment")
        elif path.suffix.lower() == ".xlsx":
            expected_label = ALLOWED_WORKBOOKS.get(relative)
            if expected_label is None:
                issues.append("workbook outside template/example locations")
            elif not _workbook_label_is_expected(path, expected_label):
                issues.append("template/example workbook content")
    return sorted(set(issues))


def main(argv: list[str] | None = None) -> int:
    arguments = argv if argv is not None else sys.argv[1:]
    if len(arguments) != 1:
        print("usage: prepublish_check.py REPOSITORY", file=sys.stderr)
        return 2
    issues = check(Path(arguments[0]))
    if issues:
        for issue in issues:
            print(f"error: prohibited {issue}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
