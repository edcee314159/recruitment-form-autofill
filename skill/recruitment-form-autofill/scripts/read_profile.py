"""Read recruitment-profile records from the approved Excel database."""

from __future__ import annotations

import argparse
import json
import sys
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Iterable
from xml.etree.ElementTree import ParseError

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException


DEFAULT_DATABASE = Path(r"D:\Codex\autumn-recruitment\personal-info-summary\秋招个人信息总表.xlsx")
GENERAL_SHEET = "资料总表"
EXPERIENCE_SHEET = "经历与成果"
GENERAL_HEADERS = (
    "序号", "分类", "字段", "当前值", "状态", "填写提示", "来源", "备注", "使用场景"
)
GENERAL_KEYS = (
    "index", "category", "field", "value", "status", "hint", "source", "note", "usage"
)
EXPERIENCE_HEADERS = (
    "类别", "项目/单位", "角色/级别", "地点", "开始/考试时间", "结束/颁发时间",
    "成绩/编号", "详细内容", "来源", "备注"
)
EXPERIENCE_KEYS = (
    "category", "organization", "role", "location", "start", "end", "result", "details", "source", "note"
)


class ProfileError(Exception):
    """A database problem that can be reported without exposing profile data."""


class SchemaError(ProfileError):
    """The workbook does not match the fixed, supported schema."""


MALFORMED_WORKBOOK_ERRORS = (OSError, zipfile.BadZipFile, InvalidFileException, ParseError)


def _rows_from_sheet(
    path: Path,
    sheet_name: str,
    header_row: int,
    headers: tuple[str, ...],
    keys: tuple[str, ...],
) -> list[dict[str, object]]:
    workbook = _open_workbook(path)
    try:
        return _rows_from_workbook(workbook, sheet_name, header_row, headers, keys)
    finally:
        workbook.close()


def _open_workbook(path: Path):
    if not path.is_file():
        raise ProfileError("database not found")
    try:
        return openpyxl.load_workbook(path, read_only=True, data_only=True)
    except MALFORMED_WORKBOOK_ERRORS as exc:
        raise ProfileError("database could not be read") from exc


def _rows_from_workbook(
    workbook,
    sheet_name: str,
    header_row: int,
    headers: tuple[str, ...],
    keys: tuple[str, ...],
) -> list[dict[str, object]]:
    try:
        if sheet_name not in workbook.sheetnames:
            raise ProfileError(f"required sheet missing: {sheet_name}")
        sheet = workbook[sheet_name]
        actual_headers = tuple(
            sheet.cell(row=header_row, column=column).value
            for column in range(1, len(headers) + 1)
        )
        trailing_headers = tuple(
            sheet.cell(row=header_row, column=column).value
            for column in range(len(headers) + 1, sheet.max_column + 1)
        )
        if actual_headers != headers or any(value is not None for value in trailing_headers):
            raise SchemaError(f"unexpected headers in sheet: {sheet_name}")
        records: list[dict[str, object]] = []
        for values in sheet.iter_rows(min_row=header_row + 1, max_col=len(headers), values_only=True):
            if not any(value is not None for value in values):
                continue
            records.append(dict(zip(keys, values)))
        return records
    except MALFORMED_WORKBOOK_ERRORS as exc:
        raise ProfileError("database could not be read") from exc


def load_general_profile(path: Path) -> list[dict[str, object]]:
    """Load general-profile rows using the fixed 资料总表 schema."""
    return _rows_from_sheet(path, GENERAL_SHEET, 8, GENERAL_HEADERS, GENERAL_KEYS)


def load_experience_records(path: Path) -> list[dict[str, object]]:
    """Load experience rows using the fixed 经历与成果 schema."""
    return _rows_from_sheet(path, EXPERIENCE_SHEET, 6, EXPERIENCE_HEADERS, EXPERIENCE_KEYS)


def load_inventory(path: Path) -> dict[str, list[object]]:
    """List profile labels using one workbook open and without returning values."""
    workbook = _open_workbook(path)
    try:
        general = _rows_from_workbook(
            workbook, GENERAL_SHEET, 8, GENERAL_HEADERS, GENERAL_KEYS
        )
        experiences = _rows_from_workbook(
            workbook, EXPERIENCE_SHEET, 6, EXPERIENCE_HEADERS, EXPERIENCE_KEYS
        )
    finally:
        workbook.close()
    return _inventory(general, experiences)


def load_bundle(
    path: Path,
    fields: Iterable[str],
    categories: Iterable[str],
) -> dict[str, list[dict[str, object]]]:
    """Read requested records from both sheets using one workbook open."""
    workbook = _open_workbook(path)
    try:
        general = _rows_from_workbook(
            workbook, GENERAL_SHEET, 8, GENERAL_HEADERS, GENERAL_KEYS
        )
        experiences = _rows_from_workbook(
            workbook, EXPERIENCE_SHEET, 6, EXPERIENCE_HEADERS, EXPERIENCE_KEYS
        )
    finally:
        workbook.close()

    selected_general: list[dict[str, object]] = []
    for field in fields:
        matches = [record for record in general if record["field"] == field]
        if not matches:
            raise ProfileError("requested field not found")
        if any(not _nonblank(record["value"]) for record in matches):
            raise ProfileError("requested field has a blank value")
        selected_general.extend(matches)

    selected_experiences: list[dict[str, object]] = []
    for category in categories:
        matches = [record for record in experiences if record["category"] == category]
        if not matches:
            raise ProfileError("requested experience category not found")
        selected_experiences.extend(matches)
    return {"general": selected_general, "experiences": selected_experiences}


def _json_default(value: object) -> str:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _write_json(payload: object) -> None:
    print(json.dumps(payload, ensure_ascii=False, default=_json_default))


def _nonblank(value: object) -> bool:
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def _inventory(general: Iterable[dict[str, object]], experiences: Iterable[dict[str, object]]) -> dict[str, list[object]]:
    return {
        "categories": sorted({record["category"] for record in general if _nonblank(record["category"])}),
        "fields": [record["field"] for record in general if _nonblank(record["field"])],
        "experience_categories": sorted(
            {record["category"] for record in experiences if _nonblank(record["category"])}
        ),
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Read the approved recruitment profile workbook.")
    parser.add_argument("--database", type=Path, default=DEFAULT_DATABASE, metavar="PATH")
    commands = parser.add_subparsers(dest="command", required=True)
    commands.add_parser("inventory", help="List profile field and category names only.")
    get_parser = commands.add_parser("get", help="Read explicitly requested general fields.")
    get_parser.add_argument("--field", action="append", required=True, metavar="FIELD")
    experiences_parser = commands.add_parser("experiences", help="Read one experience category.")
    experiences_parser.add_argument("--category", required=True, metavar="CATEGORY")
    bundle_parser = commands.add_parser("bundle", help="Read requested fields and experience categories together.")
    bundle_parser.add_argument("--field", action="append", default=[], metavar="FIELD")
    bundle_parser.add_argument("--category", action="append", default=[], metavar="CATEGORY")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        if args.command == "inventory":
            _write_json(load_inventory(args.database))
            return 0

        if args.command == "get":
            general = load_general_profile(args.database)
            selected: list[dict[str, object]] = []
            for field in args.field:
                matches = [record for record in general if record["field"] == field]
                if not matches:
                    raise ProfileError("requested field not found")
                if any(not _nonblank(record["value"]) for record in matches):
                    return _error("requested field has a blank value", 3)
                selected.extend(matches)
            _write_json({"records": selected})
            return 0

        if args.command == "bundle":
            try:
                _write_json(load_bundle(args.database, args.field, args.category))
                return 0
            except ProfileError as exc:
                if str(exc) == "requested field has a blank value":
                    return _error(str(exc), 3)
                raise

        experiences = load_experience_records(args.database)
        selected = [record for record in experiences if record["category"] == args.category]
        if not selected:
            raise ProfileError("requested experience category not found")
        _write_json({"records": selected})
        return 0
    except ProfileError as exc:
        return _error(str(exc), 2)


def _error(message: str, exit_code: int) -> int:
    print(f"error: {message}", file=sys.stderr)
    return exit_code


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
    raise SystemExit(main())
