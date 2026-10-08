"""Copy supported legacy records into a new local public-template workbook."""

from __future__ import annotations

import argparse
from dataclasses import asdict
from io import BytesIO
import json
from pathlib import Path
import sys

import openpyxl

import read_profile as reader
from profile_config import BLANK_TEMPLATE, REPOSITORY_ROOT
from validate_profile import ValidationIssue


FIELD_ALIASES = {"手机号": "手机号码", "电子邮箱": "邮箱", "毕业时间": "预计毕业时间"}


def _write_value(sheet, row, column, value):
    cell = sheet.cell(row, column, value)
    # Treat profile text as text even if it starts with a spreadsheet formula marker.
    if isinstance(value, str):
        cell.data_type = "s"


def migrate_legacy(source: Path, destination: Path) -> list[ValidationIssue]:
    """Create a new workbook and return redacted warnings for unmapped rows.

    The source and existing destinations are never written. Only the defined
    personal fields and experience columns are copied; provenance is omitted.
    """
    source = Path(source).resolve()
    destination = Path(destination).resolve()
    if destination == source or destination.exists():
        raise reader.ProfileError("migration requires a new destination")
    if destination.is_relative_to(REPOSITORY_ROOT.resolve()):
        raise reader.ProfileError("migration destination must be outside the repository")
    workbook = reader._open_workbook(source)
    warnings = []
    output = None
    try:
        if reader._detect_workbook_schema(workbook) != "legacy":
            raise reader.ProfileError("migration requires a legacy workbook")
        reader._general_from_workbook(workbook)
        experiences = reader._experiences_from_workbook(workbook)
        output = openpyxl.load_workbook(BLANK_TEMPLATE)
        mapped = set()
        for row, values in enumerate(workbook[reader.GENERAL_SHEET].iter_rows(min_row=9, max_col=len(reader.GENERAL_HEADERS), values_only=True), 9):
            if not any(reader._nonblank(v) for v in values):
                continue
            field = FIELD_ALIASES.get(values[2], values[2])
            if field not in reader.PERSONAL_HEADERS:
                warnings.append(ValidationIssue(reader.GENERAL_SHEET, row, "unmapped_field", "This legacy field has no public-template mapping and was not copied."))
                continue
            if field in mapped:
                raise reader.ProfileError("multiple legacy rows map to one personal field")
            mapped.add(field)
            _write_value(output[reader.PERSONAL_SHEET], 2, reader.PERSONAL_HEADERS.index(field) + 1, values[3])
        for row, record in enumerate(experiences, 2):
            for column, key in enumerate(reader.PUBLIC_EXPERIENCE_KEYS, 1):
                _write_value(output[reader.PUBLIC_EXPERIENCE_SHEET], row, column, record[key])
        buffer = BytesIO()
        output.save(buffer)
        destination.parent.mkdir(parents=True, exist_ok=True)
        # Exclusive creation prevents an intervening writer from being overwritten.
        with destination.open("xb") as handle:
            handle.write(buffer.getvalue())
        return warnings
    except reader.MALFORMED_WORKBOOK_ERRORS as exc:
        raise reader.ProfileError("migration could not create the destination") from exc
    finally:
        workbook.close()
        if output is not None:
            output.close()


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Migrate a legacy recruitment workbook to a new local workbook.")
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--destination", type=Path, required=True)
    args = parser.parse_args(argv)
    try:
        warnings = migrate_legacy(args.source, args.destination)
    except reader.ProfileError as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    print(json.dumps({"warnings": [asdict(issue) for issue in warnings]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
    raise SystemExit(main())
