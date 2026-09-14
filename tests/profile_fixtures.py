"""Fictional workbook fixtures; never read a real candidate profile."""

from pathlib import Path
import sys
import tempfile
import unittest

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "skill" / "recruitment-form-autofill" / "scripts"
sys.path.insert(0, str(SCRIPTS))
PERSONAL = ("姓名", "手机号码", "邮箱", "出生日期", "籍贯", "现居城市", "学校名称", "专业名称", "学历", "学位", "入学时间", "预计毕业时间", "期望岗位方向", "期望工作城市")
EXPERIENCES = ("类别", "项目/单位", "角色/级别", "地点", "开始/考试时间", "结束/颁发时间", "成绩/编号", "详细内容", "备注")


class WorkbookTestCase(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.directory = Path(self.temporary.name)
        self.path = self.directory / "fictional.xlsx"

    def public(self, rows=(), attachments=()):
        workbook = Workbook()
        workbook.active.title = "使用说明"
        sheet = workbook.create_sheet("个人资料")
        sheet.append(PERSONAL)
        sheet.append(["测试候选人（虚构）", "13800000000", "demo@example.invalid"])
        sheet = workbook.create_sheet("经历清单")
        sheet.append(EXPERIENCES)
        for row in rows:
            sheet.append(row)
        sheet = workbook.create_sheet("附件清单")
        sheet.append(["文件类型", "本地文件路径", "说明"])
        for row in attachments:
            sheet.append(row)
        workbook.save(self.path)
        workbook.close()
        return self.path

    def legacy(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "资料总表"
        for col, value in enumerate(("序号", "分类", "字段", "当前值", "状态", "填写提示", "来源", "备注", "使用场景"), 1):
            sheet.cell(8, col, value)
        sheet.append([1, "基本信息", "姓名", "测试候选人（虚构）", "确认", None, "虚构来源"])
        sheet.append([2, "联系", "手机号", "13800000000"])
        sheet.append([3, "其他", "自定义字段", "虚构保留内容"])
        sheet = workbook.create_sheet("经历与成果")
        for col, value in enumerate((*EXPERIENCES[:8], "来源", "备注"), 1):
            sheet.cell(6, col, value)
        for name, score in (("CET-4", 550), ("CET-6", 520), ("IELTS", 6.5)):
            sheet.append(["语言能力", name, None, None, "2024-06-01", None, score, "虚构成绩", "虚构来源", "虚构备注"])
        workbook.save(self.path)
        workbook.close()
        return self.path
