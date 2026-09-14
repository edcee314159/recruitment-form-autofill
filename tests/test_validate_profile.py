from dataclasses import asdict
from datetime import datetime
import importlib
import importlib.util
import json

from openpyxl import load_workbook
from tests.profile_fixtures import WorkbookTestCase


class ValidatorTest(WorkbookTestCase):
    def validate(self):
        self.assertIsNotNone(importlib.util.find_spec("validate_profile"), "validator must exist")
        return importlib.import_module("validate_profile").validate_profile(self.path)

    def test_accepts_separate_language_rows_and_valid_dates(self):
        self.public(rows=[
            ["语言能力", "CET-4", None, None, "2024-06", None, 550],
            ["语言能力", "CET-6", None, None, "2024-12-01", None, 520],
            ["语言能力", "IELTS", None, None, "2025-01-01", None, 6.5],
            ["项目经历", "虚构项目", None, None, "2025-01", "至今"],
        ])
        self.assertEqual(self.validate(), [])

    def test_checks_structure_dates_duplicates_and_attachments_without_leakage(self):
        self.public(rows=[
            ["PRIVATE-CATEGORY", "PRIVATE-ORG"],
            ["项目经历", None, "PRIVATE-ROLE"],
            ["项目经历", "PRIVATE-ORG", None, None, "2025-02-30"],
            ["项目经历", "PRIVATE-ORG", None, None, "2025-02", "2024-01"],
            ["语言能力", "CET-4", None, None, "2024-06", None, None],
            ["语言能力", "IELTS", None, None, "2024-06", None, 6.5],
            ["语言能力", "IELTS", None, None, "2024-06", None, 7],
        ], attachments=[["简历", "PRIVATE-FILE.pdf"]])
        issues = self.validate()
        self.assertTrue({"unsupported_category", "structural_blank", "invalid_date", "date_order", "language_pair", "duplicate_identity", "attachment_missing"}.issubset({issue.code for issue in issues}))
        payload = json.dumps([asdict(issue) for issue in issues], ensure_ascii=False)
        self.assertNotIn("PRIVATE", payload)
        for issue in issues:
            self.assertEqual(set(asdict(issue)), {"sheet", "row", "code", "message"})
        self.assertIn(("经历清单", 8, "duplicate_identity"), [(i.sheet, i.row, i.code) for i in issues])

    def test_header_and_extra_personal_rows_are_reported(self):
        self.public()
        workbook = load_workbook(self.path)
        workbook["经历清单"]["B1"] = "PRIVATE-HEADER"
        workbook["个人资料"]["A3"] = "PRIVATE-NAME"
        workbook["个人资料"]["D2"] = "2024/13/88"
        workbook.save(self.path)
        workbook.close()
        issues = self.validate()
        self.assertTrue({"headers", "extra_personal_row", "invalid_date"}.issubset({i.code for i in issues}))
        self.assertNotIn("PRIVATE", repr(issues))

    def test_relative_attachment_resolves_beside_workbook(self):
        attachment = self.directory / "fictional.pdf"
        attachment.write_bytes(b"fictional")
        self.public(attachments=[["简历", "fictional.pdf"]])
        self.assertEqual(self.validate(), [])

    def test_legacy_can_be_validated_without_migration(self):
        self.legacy()
        self.assertEqual(self.validate(), [])

    def test_personal_date_order_and_combined_language_scores_are_rejected(self):
        self.public(rows=[["语言能力", "CET-4 / CET-6", None, None, None, None, "550 / 520"]])
        workbook = load_workbook(self.path)
        workbook["个人资料"]["K2"] = "2025-09"
        workbook["个人资料"]["L2"] = "2024-06"
        workbook.save(self.path)
        workbook.close()
        self.assertTrue({"date_order", "language_pair"}.issubset({i.code for i in self.validate()}))

    def test_accepts_bilingual_aliases_for_one_language_credential(self):
        self.public(rows=[
            ["语言能力", "CET-4（大学英语四级）", None, None, "2024-06", None, 550],
            ["语言能力", "IELTS（雅思）", None, None, "2025-01", None, 6.5],
        ])
        self.assertEqual(self.validate(), [])

    def test_detects_duplicates_when_equivalent_dates_use_different_cell_types(self):
        self.public(rows=[
            ["语言能力", "IELTS", None, None, datetime(2024, 6, 1), None, 6.5],
            ["语言能力", "IELTS", None, None, "2024-06-01", None, 6.5],
        ])
        issues = self.validate()
        self.assertIn(("经历清单", 3, "duplicate_identity"), [(i.sheet, i.row, i.code) for i in issues])
