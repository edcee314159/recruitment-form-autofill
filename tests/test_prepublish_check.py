from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
CHECKER = ROOT / "scripts" / "prepublish_check.py"


class PrepublishCheckTest(unittest.TestCase):
    def test_rejects_filled_template_despite_unchanged_blank_label(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            target = root / "templates" / "秋招个人资料库模板.xlsx"
            target.parent.mkdir()
            workbook = load_workbook(ROOT / "templates" / target.name)
            workbook["个人资料"]["A2"] = "PRIVATE-FICTIONAL-CANDIDATE"
            workbook.save(target)
            workbook.close()
            result = self.run_check(root)
        self.assertEqual(result.returncode, 2)
        self.assertNotIn("PRIVATE-FICTIONAL-CANDIDATE", result.stderr)

    def test_rejects_modified_example_with_unchanged_fictional_label(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            target = root / "examples" / "示例资料库.xlsx"
            target.parent.mkdir()
            workbook = load_workbook(ROOT / "examples" / target.name)
            workbook["个人资料"]["A2"] = "PRIVATE-FICTIONAL-CANDIDATE"
            workbook.save(target)
            workbook.close()
            result = self.run_check(root)
        self.assertEqual(result.returncode, 2)
        self.assertNotIn("PRIVATE-FICTIONAL-CANDIDATE", result.stderr)

    def run_check(self, root):
        return subprocess.run(
            [sys.executable, str(CHECKER), str(root)],
            cwd=ROOT,
            capture_output=True,
            text=True,
            check=False,
        )

    def test_rejects_personal_workbook_without_printing_contents(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            (root / "candidate.personal.xlsx").write_bytes(b"PRIVATE-WORKBOOK-CONTENT")
            result = self.run_check(root)
        self.assertEqual(result.returncode, 2)
        self.assertIn("personal workbook", result.stderr)
        self.assertNotIn("PRIVATE-WORKBOOK-CONTENT", result.stderr)

    def test_allows_only_blank_or_fictional_xlsx_in_the_example_locations(self):
        result = self.run_check(ROOT)
        self.assertEqual(result.returncode, 0, result.stderr)
