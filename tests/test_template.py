from pathlib import Path
import os
import shutil
import subprocess
import tempfile
import unittest

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "templates" / "秋招个人资料库模板.xlsx"
EXAMPLE = ROOT / "examples" / "示例资料库.xlsx"

PERSONAL_HEADERS = [
    "姓名", "手机号码", "邮箱", "出生日期", "籍贯", "现居城市", "学校名称", "专业名称",
    "学历", "学位", "入学时间", "预计毕业时间", "期望岗位方向", "期望工作城市",
]
EXPERIENCE_HEADERS = [
    "类别", "项目/单位", "角色/级别", "地点", "开始/考试时间", "结束/颁发时间",
    "成绩/编号", "详细内容", "备注",
]
ATTACHMENT_HEADERS = ["文件类型", "本地文件路径", "说明"]
EXPERIENCE_CATEGORIES = [
    "教育经历", "实习经历", "工作经历", "项目经历", "语言能力", "专业证书",
    "竞赛获奖", "荣誉称号", "校园活动", "科研成果",
]


def find_node_executable():
    override = os.environ.get("CODEX_NODE")
    if override:
        executable = Path(override).expanduser()
        if executable.is_file():
            return str(executable)
        raise RuntimeError(
            f"CODEX_NODE points to a missing Node.js executable: {executable}"
        )

    executable = shutil.which("node")
    if executable:
        return executable
    raise RuntimeError(
        "Node.js was not found. Install Node.js and add 'node' to PATH, "
        "or set CODEX_NODE to the Node.js executable path."
    )


NODE = find_node_executable()


class TemplateWorkbookTest(unittest.TestCase):
    def generate(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        output_root = Path(temporary.name)
        subprocess.run(
            [str(NODE), "scripts/create_template.mjs", "--output-dir", str(output_root)],
            cwd=ROOT,
            check=True,
        )
        return (
            output_root / "templates" / "秋招个人资料库模板.xlsx",
            output_root / "examples" / "示例资料库.xlsx",
        )

    def test_generator_creates_blank_template_with_required_structure(self):
        """Catches a generator that omits required sheets, headers, or blank rows."""
        template, _ = self.generate()

        workbook = load_workbook(template)
        self.assertEqual(workbook.sheetnames, ["使用说明", "个人资料", "经历清单", "附件清单"])
        self.assertEqual(workbook["个人资料"]["A1"].value, "姓名")
        self.assertEqual(workbook["经历清单"]["A1"].value, "类别")
        self.assertEqual(
            [cell.value for cell in workbook["个人资料"][1]], PERSONAL_HEADERS
        )
        self.assertEqual(
            [cell.value for cell in workbook["经历清单"][1]], EXPERIENCE_HEADERS
        )
        self.assertEqual(
            [cell.value for cell in workbook["附件清单"][1]], ATTACHMENT_HEADERS
        )
        self.assertTrue(
            all(cell.value is None for cell in workbook["个人资料"][2])
        )
        self.assertTrue(
            all(
                all(cell.value is None for cell in row)
                for row in workbook["经历清单"].iter_rows(min_row=2)
            )
        )
        self.assertTrue(
            all(
                all(cell.value is None for cell in row)
                for row in workbook["附件清单"].iter_rows(min_row=2)
            )
        )

    def test_generator_creates_fictional_example_with_language_records(self):
        """Catches an example that loses the required separate language rows."""
        _, example = self.generate()

        workbook = load_workbook(example)
        experience_sheet = workbook["经历清单"]
        categories = [
            row[0].value for row in experience_sheet.iter_rows(min_row=2, values_only=False)
        ]
        language_rows = [
            row for row in experience_sheet.iter_rows(min_row=2, values_only=True)
            if row[0] == "语言能力"
        ]
        language_names = {row[1] for row in language_rows}
        self.assertIn("语言能力", categories)
        self.assertEqual(language_names, {"CET-4", "CET-6", "IELTS"})
        self.assertEqual(len(language_rows), 3)

    def test_template_limits_experience_categories_to_the_published_list(self):
        """Catches a template that leaves experience categories unconstrained."""
        template, _ = self.generate()

        workbook = load_workbook(template)
        validations = workbook["经历清单"].data_validations.dataValidation
        category_validation = next(
            validation
            for validation in validations
            if "A2:A501" in str(validation.sqref)
        )
        self.assertEqual(category_validation.type, "list")
        self.assertEqual(
            category_validation.formula1.strip('"').split(","), EXPERIENCE_CATEGORIES
        )

    def test_example_experience_role_column_fits_the_program_name(self):
        """Catches a narrow role column that clips the example program name."""
        _, example = self.generate()

        workbook = load_workbook(example)
        role_column = workbook["经历清单"].column_dimensions["C"]
        self.assertGreaterEqual(role_column.width, 30)
