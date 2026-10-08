import importlib
import importlib.util
import json
import subprocess
import sys
from openpyxl import load_workbook

from tests.profile_fixtures import WorkbookTestCase, ROOT
import read_profile


class MigrationTest(WorkbookTestCase):
    def test_documented_cli_migrates_records_and_reports_redacted_warnings(self):
        self.legacy()
        original = self.path.read_bytes()
        destination = self.directory / "cli-result.xlsx"
        result = subprocess.run([sys.executable, "-B", str(ROOT / "skill/recruitment-form-autofill/scripts/migrate_profile.py"), "--source", str(self.path), "--destination", str(destination)], capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertTrue(destination.is_file(), "CLI must create the migrated workbook")
        self.assertEqual(self.path.read_bytes(), original)
        self.assertEqual(read_profile.load_general_profile(destination)[0]["value"], "测试候选人（虚构）")
        self.assertEqual(json.loads(result.stdout)["warnings"][0]["code"], "unmapped_field")
        self.assertNotIn("虚构保留内容", result.stdout)

    def test_cli_rejects_missing_source_instead_of_succeeding_silently(self):
        destination = self.directory / "cli-result.xlsx"
        result = subprocess.run([sys.executable, "-B", str(ROOT / "skill/recruitment-form-autofill/scripts/migrate_profile.py"), "--source", str(self.directory / "PRIVATE-MISSING.xlsx"), "--destination", str(destination)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 2)
        self.assertFalse(destination.exists())
        self.assertNotIn("PRIVATE-MISSING", result.stderr)

    def migration(self):
        self.assertIsNotNone(importlib.util.find_spec("migrate_profile"), "migration must exist")
        return importlib.import_module("migrate_profile")

    def test_migration_preserves_source_and_copies_only_public_fields(self):
        self.legacy()
        original = self.path.read_bytes()
        destination = self.directory / "public.xlsx"
        warnings = self.migration().migrate_legacy(self.path, destination)
        self.assertEqual(self.path.read_bytes(), original)
        self.assertEqual(read_profile.detect_schema(destination), "public")
        general = read_profile.load_bundle(destination, ["姓名", "手机号码"], [])
        self.assertEqual([r["value"] for r in general["general"]], ["测试候选人（虚构）", "13800000000"])
        experiences = read_profile.load_experience_records(destination)
        self.assertEqual([(r["organization"], r["result"]) for r in experiences], [("CET-4", 550), ("CET-6", 520), ("IELTS", 6.5)])
        workbook = load_workbook(destination)
        self.assertEqual(workbook.sheetnames, ["使用说明", "个人资料", "经历清单", "附件清单"])
        values = repr([tuple(row) for sheet in workbook for row in sheet.iter_rows(values_only=True)])
        workbook.close()
        self.assertNotIn("虚构保留内容", values)
        self.assertNotIn("虚构来源", values)
        self.assertEqual([(w.sheet, w.row, w.code) for w in warnings], [("资料总表", 11, "unmapped_field")])
        self.assertNotIn("虚构保留内容", repr(warnings))

    def test_existing_destination_and_source_alias_are_never_overwritten(self):
        self.legacy()
        module = self.migration()
        original = self.path.read_bytes()
        with self.assertRaises(read_profile.ProfileError):
            module.migrate_legacy(self.path, self.path)
        destination = self.directory / "existing.xlsx"
        destination.write_bytes(b"keep existing content")
        with self.assertRaises(read_profile.ProfileError):
            module.migrate_legacy(self.path, destination)
        self.assertEqual(destination.read_bytes(), b"keep existing content")
        self.assertEqual(self.path.read_bytes(), original)

    def test_public_input_is_rejected_without_creating_output(self):
        self.public()
        destination = self.directory / "public.xlsx"
        with self.assertRaises(read_profile.ProfileError):
            self.migration().migrate_legacy(self.path, destination)
        self.assertFalse(destination.exists())

    def test_migration_rejects_repository_destination(self):
        self.legacy()
        with self.assertRaises(read_profile.ProfileError):
            self.migration().migrate_legacy(self.path, ROOT / "private-migration-test.xlsx")
        self.assertFalse((ROOT / "private-migration-test.xlsx").exists())

    def test_literal_formula_like_text_is_not_turned_into_a_formula(self):
        self.legacy()
        workbook = load_workbook(self.path)
        cell = workbook["资料总表"]["D9"]
        cell.value = "=FICTIONAL(1)"
        cell.data_type = "s"
        workbook.save(self.path)
        workbook.close()
        destination = self.directory / "public.xlsx"
        self.migration().migrate_legacy(self.path, destination)
        workbook = load_workbook(destination, data_only=False)
        self.assertEqual(workbook["个人资料"]["A2"].value, "=FICTIONAL(1)")
        self.assertEqual(workbook["个人资料"]["A2"].data_type, "s")
        workbook.close()
